"""
Generate the 10-slide Knowledge Assistant API presentation.
Outputs Docs/presentation.pptx (dark navy theme).

Usage:
    pip install python-pptx
    python generate_presentation.py
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Palette ──────────────────────────────────────────────────────────────────
NAVY  = RGBColor(0x0D, 0x1B, 0x2A)
CYAN  = RGBColor(0x00, 0xC6, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY  = RGBColor(0xCC, 0xCC, 0xCC)
GREEN = RGBColor(0x00, 0xE5, 0x96)
AMBER = RGBColor(0xFF, 0xC1, 0x07)

W = Inches(13.33)
H = Inches(7.5)


# ── Helpers ───────────────────────────────────────────────────────────────────
def bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = NAVY


def rect(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def txt(slide, text, x, y, w, h, size=24, bold=False, color=WHITE, align=PP_ALIGN.LEFT, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.italic = italic
    return tb


def multiline_txt(slide, lines, x, y, w, h, size=18, color=WHITE):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.color.rgb = color


def pill(slide, label, x, y, color=CYAN):
    r = rect(slide, x, y, 2.8, 0.45, color)
    txt(slide, label, x + 0.05, y + 0.05, 2.7, 0.35, size=14, bold=True,
        color=NAVY, align=PP_ALIGN.CENTER)


def accent_bar(slide, y=0.0, h=0.08):
    rect(slide, 0, y, 13.33, h, CYAN)


# ── Slides ────────────────────────────────────────────────────────────────────

def slide_01_title(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    accent_bar(slide, y=0.0, h=0.12)
    accent_bar(slide, y=7.38, h=0.12)
    rect(slide, 0.6, 1.6, 12.0, 0.06, CYAN)
    txt(slide, "Knowledge Assistant API", 0.6, 1.8, 12.0, 1.2,
        size=52, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    txt(slide, "RAG + FastAPI + LangChain + TruLens",
        0.6, 3.1, 12.0, 0.7, size=26, color=CYAN, align=PP_ALIGN.CENTER)
    txt(slide, "Build, manage and evaluate production-grade retrieval-augmented generation systems",
        0.6, 3.9, 12.0, 0.8, size=18, color=GRAY, align=PP_ALIGN.CENTER, italic=True)
    for i, label in enumerate(["FastAPI", "LangChain", "OpenAI", "ChromaDB / FAISS", "TruLens"]):
        pill(slide, label, 1.0 + i * 2.4, 5.6)


def slide_02_problem(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    accent_bar(slide)
    txt(slide, "The Problem", 0.6, 0.2, 12.0, 0.8, size=36, bold=True, color=CYAN)
    points = [
        ("LLMs hallucinate", "Without grounding, models invent facts not in their training data."),
        ("Knowledge is stale", "Pre-trained models cannot see your private or recent documents."),
        ("No accountability", "How do you know if a RAG answer is actually correct?"),
        ("Hard to manage docs", "Uploading, updating, and versioning documents lacks tooling."),
    ]
    for i, (title, desc) in enumerate(points):
        y = 1.3 + i * 1.3
        rect(slide, 0.5, y, 12.2, 1.1, RGBColor(0x12, 0x26, 0x3D))
        txt(slide, title, 0.8, y + 0.08, 4.0, 0.5, size=20, bold=True, color=AMBER)
        txt(slide, desc, 0.8, y + 0.55, 11.5, 0.45, size=16, color=GRAY)


def slide_03_solution(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    accent_bar(slide)
    txt(slide, "Our Solution", 0.6, 0.2, 12.0, 0.8, size=36, bold=True, color=CYAN)
    txt(slide, "A fully REST-ful Knowledge Assistant with built-in quality evaluation",
        0.6, 0.95, 12.0, 0.5, size=18, color=GRAY, italic=True)
    steps = [
        ("1  Upload", "Push .txt files into isolated knowledge bases via REST"),
        ("2  Embed", "LangChain chunks & embeds text using OpenAI ada-002"),
        ("3  Store", "FAISS (speed) or ChromaDB (persistence) holds vectors"),
        ("4  Retrieve", "Semantic similarity search returns top-5 chunks"),
        ("5  Generate", "GPT synthesises a grounded answer from retrieved context"),
        ("6  Evaluate", "TruLens RAG Triad scores quality automatically"),
    ]
    for i, (label, desc) in enumerate(steps):
        col, row = divmod(i, 3)
        x = 0.5 + col * 6.5
        y = 1.6 + row * 1.8
        rect(slide, x, y, 6.0, 1.55, RGBColor(0x12, 0x26, 0x3D))
        txt(slide, label, x + 0.2, y + 0.1, 5.6, 0.5, size=20, bold=True, color=GREEN)
        txt(slide, desc, x + 0.2, y + 0.65, 5.6, 0.7, size=15, color=GRAY)


def slide_04_arch(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    accent_bar(slide)
    txt(slide, "System Architecture", 0.6, 0.2, 12.0, 0.8, size=36, bold=True, color=CYAN)
    layers = [
        ("Client Layer",        "HTTP / Swagger UI / Jupyter Notebook"),
        ("API Layer",           "FastAPI  |  Pydantic v2 schemas  |  15 endpoints"),
        ("Business Logic",      "kb_manager.py  |  rag.py  |  trulens_eval.py"),
        ("LangChain Pipeline",  "TextLoader  →  Splitter  →  Embeddings  →  Retriever  →  LLM"),
        ("Storage Layer",       "data/{kb_id}/documents/   +   vector_store/   +   metadata.json"),
    ]
    colors = [RGBColor(0x00, 0x5F, 0x73), RGBColor(0x00, 0x7A, 0x8A),
              RGBColor(0x0A, 0x95, 0x9B), RGBColor(0x16, 0xAF, 0xA0), RGBColor(0x22, 0xC5, 0x9A)]
    for i, ((layer, detail), color) in enumerate(zip(layers, colors)):
        y = 1.3 + i * 1.1
        rect(slide, 0.5, y, 12.2, 0.95, color)
        txt(slide, layer, 0.8, y + 0.1, 3.0, 0.5, size=17, bold=True, color=NAVY)
        txt(slide, detail, 3.9, y + 0.15, 8.5, 0.5, size=15, color=NAVY)


def slide_05_endpoints(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    accent_bar(slide)
    txt(slide, "REST API Endpoints (15 total)", 0.6, 0.2, 12.0, 0.8, size=36, bold=True, color=CYAN)
    groups = [
        ("System",          ["GET  /health"]),
        ("Knowledge Bases", ["POST /knowledge-bases/", "GET  /knowledge-bases/",
                             "GET  /knowledge-bases/{id}", "PUT  /knowledge-bases/{id}",
                             "DEL  /knowledge-bases/{id}"]),
        ("Documents",       ["POST /…/documents/upload", "GET  /…/documents/",
                             "GET  /…/documents/{name}", "PUT  /…/documents/{name}",
                             "DEL  /…/documents/{name}"]),
        ("Query",           ["POST /…/query/"]),
        ("Evaluation",      ["POST /…/evaluate/", "GET  /evaluate/records",
                             "GET  /evaluate/leaderboard"]),
    ]
    col_w = 2.5
    for col, (tag, routes) in enumerate(groups):
        x = 0.4 + col * 2.55
        txt(slide, tag, x, 1.15, col_w, 0.4, size=14, bold=True, color=CYAN)
        rect(slide, x, 1.5, col_w - 0.1, 0.04, CYAN)
        for row, route in enumerate(routes):
            method, path = route.split(None, 1)
            m_color = GREEN if method in ("GET",) else AMBER if method == "POST" else RGBColor(0xFF, 0x63, 0x47)
            txt(slide, method, x, 1.65 + row * 0.95, 0.7, 0.45, size=11, bold=True, color=m_color)
            txt(slide, path, x + 0.72, 1.65 + row * 0.95, col_w - 0.75, 0.9, size=11, color=GRAY)


def slide_06_langchain(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    accent_bar(slide)
    txt(slide, "LangChain RAG Pipeline", 0.6, 0.2, 12.0, 0.8, size=36, bold=True, color=CYAN)
    steps = [
        ("TextLoader", "Reads .txt files from disk\nwith UTF-8 encoding"),
        ("Text Splitter", "RecursiveCharacterTextSplitter\nchunk_size=500, overlap=50"),
        ("Embeddings", "OpenAIEmbeddings\ntext-embedding-ada-002\n1536 dimensions"),
        ("Vector Store", "FAISS (in-memory)\nor ChromaDB (disk)"),
        ("Retriever", "similarity search\ntop-k = 5 chunks"),
        ("LLM", "OpenAI GPT-3.5 / GPT-4\ngenerates grounded answer"),
    ]
    for i, (label, detail) in enumerate(steps):
        x = 0.3 + i * 2.15
        rect(slide, x, 1.4, 1.95, 2.4, RGBColor(0x0A, 0x2A, 0x4A))
        txt(slide, label, x + 0.1, 1.5, 1.75, 0.55, size=14, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
        txt(slide, detail, x + 0.1, 2.1, 1.75, 1.55, size=12, color=GRAY, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            txt(slide, "->", x + 1.95, 2.3, 0.25, 0.4, size=20, bold=True, color=AMBER, align=PP_ALIGN.CENTER)
    txt(slide, "Each query rebuilds the vector index on-the-fly from all .txt files in the knowledge base.",
        0.6, 4.05, 12.0, 0.5, size=15, color=GRAY, italic=True, align=PP_ALIGN.CENTER)
    txt(slide, "Vector store choice (FAISS vs ChromaDB) is set at KB creation time and cannot be changed.",
        0.6, 4.5, 12.0, 0.5, size=15, color=GRAY, italic=True, align=PP_ALIGN.CENTER)


def slide_07_trulens(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    accent_bar(slide)
    txt(slide, "TruLens RAG Triad Evaluation", 0.6, 0.2, 12.0, 0.8, size=36, bold=True, color=CYAN)
    metrics = [
        ("Answer\nRelevance", "Is the response relevant\nto the question?",
         "LLM judges response vs. query", GREEN),
        ("Context\nRelevance", "Is the retrieved context\nrelevant to the query?",
         "LLM judges each chunk vs. query\nthen averages the scores", AMBER),
        ("Groundedness", "Is the response supported\nby the retrieved context?",
         "LLM checks claims in response\nagainst source chunks (CoT)", RGBColor(0xFF, 0x63, 0x47)),
    ]
    for i, (name, question, how, color) in enumerate(metrics):
        x = 0.5 + i * 4.2
        rect(slide, x, 1.25, 3.9, 4.2, RGBColor(0x10, 0x22, 0x38))
        rect(slide, x, 1.25, 3.9, 0.06, color)
        txt(slide, name, x + 0.15, 1.4, 3.6, 0.9, size=22, bold=True, color=color, align=PP_ALIGN.CENTER)
        txt(slide, question, x + 0.15, 2.4, 3.6, 0.9, size=15, color=WHITE, align=PP_ALIGN.CENTER)
        txt(slide, how, x + 0.15, 3.4, 3.6, 1.5, size=13, color=GRAY, align=PP_ALIGN.CENTER, italic=True)
        txt(slide, "0 to 1", x + 0.15, 4.9, 3.6, 0.4, size=14, bold=True, color=color, align=PP_ALIGN.CENTER)
    txt(slide, "rag_triad_avg = mean(answer_relevance, context_relevance, groundedness)",
        0.6, 5.65, 12.0, 0.5, size=16, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
    txt(slide, "All scoring is done by an OpenAI LLM-as-judge — results are stored in eval_records.json",
        0.6, 6.15, 12.0, 0.4, size=13, color=GRAY, italic=True, align=PP_ALIGN.CENTER)


def slide_08_schema(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    accent_bar(slide)
    txt(slide, "Pydantic v2 Schemas", 0.6, 0.2, 12.0, 0.8, size=36, bold=True, color=CYAN)
    schemas = [
        ("KnowledgeBaseCreate", "name, description, vector_store"),
        ("KnowledgeBase",       "+ kb_id, created_at, updated_at, document_count"),
        ("DocumentInfo",        "name, size_bytes, created_at, updated_at"),
        ("QueryRequest/Response", "query → kb_id, query, response"),
        ("EvaluateRequest",     "query"),
        ("EvalRecord",          "id, timestamp, kb_id, query, response, scores"),
        ("EvalScores",          "answer_relevance, context_relevance,\ngroundedness, rag_triad_avg"),
        ("LeaderboardResponse", "total_evaluations, kb_id, averages"),
    ]
    for i, (name, fields) in enumerate(schemas):
        col, row = divmod(i, 4)
        x = 0.4 + col * 6.4
        y = 1.2 + row * 1.45
        rect(slide, x, y, 6.1, 1.3, RGBColor(0x10, 0x22, 0x38))
        txt(slide, name, x + 0.15, y + 0.1, 5.8, 0.5, size=16, bold=True, color=CYAN)
        txt(slide, fields, x + 0.15, y + 0.6, 5.8, 0.55, size=13, color=GRAY)


def slide_09_demo(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    accent_bar(slide)
    txt(slide, "Live Demo Flow", 0.6, 0.2, 12.0, 0.8, size=36, bold=True, color=CYAN)
    steps = [
        ("1", "Start server",   "uvicorn main:app --reload  →  http://localhost:8000/docs"),
        ("2", "Create KB",      "POST /knowledge-bases/  {name, description, vector_store}"),
        ("3", "Upload doc",     "POST /…/documents/upload  (multipart .txt file)"),
        ("4", "Ask question",   "POST /…/query/  {query: 'What is RAG?'}"),
        ("5", "Evaluate",       "POST /…/evaluate/  →  returns RAG Triad scores 0-1"),
        ("6", "Leaderboard",    "GET  /evaluate/leaderboard  →  aggregated averages"),
    ]
    for i, (num, title, cmd) in enumerate(steps):
        y = 1.35 + i * 0.98
        rect(slide, 0.4, y, 0.7, 0.75, CYAN)
        txt(slide, num, 0.4, y + 0.1, 0.7, 0.55, size=22, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        txt(slide, title, 1.25, y + 0.05, 2.5, 0.45, size=17, bold=True, color=WHITE)
        rect(slide, 3.9, y, 9.0, 0.75, RGBColor(0x08, 0x18, 0x28))
        txt(slide, cmd, 4.05, y + 0.15, 8.7, 0.45, size=13, color=GREEN)


def slide_10_summary(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    accent_bar(slide)
    accent_bar(slide, y=7.38, h=0.12)
    txt(slide, "Summary & Key Takeaways", 0.6, 0.2, 12.0, 0.8, size=36, bold=True, color=CYAN)
    points = [
        "Production-ready RAG system with full CRUD for KBs and documents",
        "Dual vector store support: FAISS for speed, ChromaDB for persistence",
        "TruLens RAG Triad gives objective, LLM-based quality scores",
        "Pydantic v2 schemas enforce request/response contracts across all 15 endpoints",
        "Jupyter notebook enables interactive TruLens experiments",
        "Architecture scales from a single KB to hundreds with isolated namespaces",
    ]
    for i, pt in enumerate(points):
        y = 1.3 + i * 0.9
        rect(slide, 0.5, y + 0.2, 0.4, 0.4, CYAN)
        txt(slide, "+", 0.5, y + 0.1, 0.4, 0.5, size=22, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        txt(slide, pt, 1.1, y + 0.1, 11.5, 0.65, size=17, color=WHITE)
    txt(slide, "github.com/anarojoecheburua/RAG-with-Langchain-and-FastAPI",
        0.6, 7.0, 12.0, 0.35, size=14, color=CYAN, align=PP_ALIGN.CENTER)


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    prs = Presentation()
    prs.slide_width  = W
    prs.slide_height = H

    slide_01_title(prs)
    slide_02_problem(prs)
    slide_03_solution(prs)
    slide_04_arch(prs)
    slide_05_endpoints(prs)
    slide_06_langchain(prs)
    slide_07_trulens(prs)
    slide_08_schema(prs)
    slide_09_demo(prs)
    slide_10_summary(prs)

    out = Path("Docs") / "presentation.pptx"
    out.parent.mkdir(exist_ok=True)
    prs.save(str(out))
    print(f"Saved {out}  ({out.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
